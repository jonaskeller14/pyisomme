# python-pptx ships no type information: `slides`/`shapes`/`placeholders` hang off a
# `lazyproperty` descriptor no checker can see through.
# pyright: reportAttributeAccessIssue=false, reportIndexIssue=false
from __future__ import annotations
from pptx.presentation import Presentation

from pyisomme.report.base_report import BaseReport
from pyisomme.report.page.base import Page


class Page_Cover(Page[BaseReport]):
    """Title slide: the report's title over its name and the tests it covers."""

    name = "Cover"
    title: str
    subtitle: str

    def __init__(self, report: BaseReport) -> None:
        super().__init__(report)
        self.title = report.title
        labels = " | ".join(report.coverage_labels)
        self.subtitle = f"{report.name}\n{labels}" if labels else report.name

    def construct(self, presentation: Presentation) -> None:
        title_slide_layout = presentation.slide_layouts[0]
        slide = presentation.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = self.title
        slide.placeholders[1].text = self.subtitle
