# python-pptx ships no type information: `slides`/`shapes`/`placeholders` hang off a
# `lazyproperty` descriptor no checker can see through.
# pyright: reportAttributeAccessIssue=false, reportIndexIssue=false
from __future__ import annotations

import io
from abc import ABC, abstractmethod

import matplotlib.pyplot as plt
from matplotlib.figure import Figure
from pptx.presentation import Presentation
from typing_extensions import override

from pyisomme.report.page.content import Page_Content

FIGSIZE_Y = 8.0


class Page_Figure(Page_Content, ABC):
    """
    A content slide whose body is a single matplotlib figure.

    Measuring the layout's content placeholder, removing it and rendering into
    the freed area is the same for every figure page, so subclasses only
    implement figure().
    """

    @override
    def construct_pptx(self, presentation: Presentation) -> None:
        super().construct_pptx(presentation)
        slide = presentation.slides[-1]

        placeholder = slide.placeholders[1]
        top = placeholder.top
        left = placeholder.left
        height = placeholder.height
        width = placeholder.width

        element = placeholder.element
        element.getparent().remove(element)

        fig = self.figure((FIGSIZE_Y * float(width) / float(height), FIGSIZE_Y))

        image_stream = io.BytesIO()
        fig.savefig(image_stream, transparent=True, bbox_inches="tight")
        # Reports build dozens of figures; without this pyplot keeps every one of
        # them alive and warns past 20.
        plt.close(fig)
        slide.shapes.add_picture(image_stream, left=left, top=top, height=height)

    @abstractmethod
    def figure(self, figsize: tuple[float, float]) -> Figure:
        """Build the figure filling the slide's content area."""
