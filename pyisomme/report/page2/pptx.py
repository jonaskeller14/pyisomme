from __future__ import annotations

import getpass
import io
import os
from dataclasses import dataclass
from datetime import datetime
from typing import Any

import plotly.graph_objects as go
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Inches


def _current_user() -> str:
    try:
        return os.getlogin()
    except OSError:
        try:
            return getpass.getuser()
        except Exception:
            return "unknown"


def default_footer() -> str:
    return f"{datetime.now().strftime('%d.%m.%Y')} | {_current_user()}"


@dataclass(frozen=True)
class ContentSlide:
    """A standard content slide and its still-empty body placeholder."""

    slide: Slide
    body_placeholder: Any


def add_content_slide(
    presentation: Presentation, *, title: str, footer: str | None
) -> ContentSlide:
    """Create the legacy content layout including its title and footer."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[1]) # pyright: ignore[reportAttributeAccessIssue]
    slide.shapes.title.text = title

    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    text_box = slide.shapes.add_textbox(
        0, Inches(slide_height / Inches(1) - 0.3), slide_width, Inches(0.3) # pyright: ignore[reportOptionalOperand]
    )
    text_frame = text_box.text_frame
    text_frame.margin_top = Inches(0.05)
    text_frame.margin_bottom = Inches(0.05)
    text_frame.margin_left = Inches(0.05)
    text_frame.margin_right = Inches(0.05)
    paragraph = text_frame.paragraphs[0]
    paragraph.text = default_footer() if footer is None else footer
    paragraph.font.size = Inches(0.2)

    return ContentSlide(slide=slide, body_placeholder=slide.placeholders[1])


def add_figure(content_slide: ContentSlide, figure: go.Figure) -> None:
    """Replace the content placeholder with a static rendering of a Plotly figure."""
    placeholder = content_slide.body_placeholder
    top = placeholder.top
    left = placeholder.left
    height = placeholder.height
    element = placeholder.element
    element.getparent().remove(element)

    image = figure.to_image(format="png", scale=2)

    content_slide.slide.shapes.add_picture( # pyright: ignore[reportAttributeAccessIssue]
        io.BytesIO(image), left=left, top=top, height=height
    )
