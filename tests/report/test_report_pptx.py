from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from pyisomme.report import ALL_REPORTS
from tests.report.report_factory import REPORT_FACTORIES


@pytest.mark.pptx
class TestReportPptx:
    @pytest.mark.parametrize("report_cls", ALL_REPORTS, ids=lambda cls: cls.__name__)
    def test_requested_reports(
        self, pytestconfig: pytest.Config, report_cls: type
    ) -> None:
        pptx_report = pytestconfig.getoption("--report")
        if pptx_report is not None and report_cls.__name__ != pptx_report.strip():
            pytest.skip(f"Not selected by --report={pptx_report}")

        output_dir = Path("out")
        output_dir.mkdir(exist_ok=True)

        path = output_dir / f"{report_cls.__name__}.pptx"
        old_modified_ns = path.stat().st_mtime_ns if path.is_file() else 0

        report = REPORT_FACTORIES[report_cls]()

        report.calculate()
        report.export_pptx(path=path)

        presentation = Presentation(path)
        assert len(presentation.slides) >= 1
        assert len(presentation.slides) == len(report.selected_pages)
        assert path.is_file()
        assert path.stat().st_mtime_ns > old_modified_ns
